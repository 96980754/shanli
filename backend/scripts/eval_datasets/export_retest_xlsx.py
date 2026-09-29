#!/usr/bin/env python3
"""把「甲方评价为勉强可接受/不可接受」的题的补测结果导出成 Excel。

输入：
  - 补测跑批产物 synthetic/agent_e2e_<name>.jsonl（含题目、甲方标准答案、系统答案）
  - 补测评分报告 reports/citation_accuracy_<name>.json（含前 5 条引用条目与 judge 判定）
  - 甲方评价 aip_ratings.jsonl（跑批产物不带评价列，按 query 关联）

用法：
    python /app/scripts/eval_datasets/export_retest_xlsx.py \
        --results /app/scripts/eval_datasets/synthetic/agent_e2e_retest_aip_20260928.jsonl \
        --report  /app/scripts/eval_datasets/reports/citation_accuracy_retest_aip_20260928.json \
        --ratings /app/scripts/eval_datasets/aip_ratings.jsonl \
        --output  /app/scripts/eval_datasets/reports/引用准确率补测结果.xlsx
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    ("序号", 6),
    ("分区", 10),
    ("问题", 46),
    ("甲方评价", 12),
    ("甲方标准答案", 46),
    ("本次补测系统答案", 70),
    ("前 5 位引用条目（文档 · chunk）", 46),
    ("引用是否正确", 12),
    ("判定说明", 50),
]

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
HIT_FILL = PatternFill("solid", fgColor="E2EFDA")
MISS_FILL = PatternFill("solid", fgColor="FCE4E4")
NA_FILL = PatternFill("solid", fgColor="F2F2F2")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="补测结果导出 Excel")
    parser.add_argument("--results", required=True, help="补测跑批产物 JSONL")
    parser.add_argument("--report", required=True, help="补测评分报告 JSON")
    parser.add_argument("--ratings", required=True, help="每行 {query, aip_rating} 的 JSONL")
    parser.add_argument("--output", required=True, help="输出 xlsx 路径")
    return parser.parse_args()


def load_ratings(path: str) -> dict[str, str]:
    ratings = {}
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        if line.strip():
            record = json.loads(line)
            ratings[record["query"]] = record.get("aip_rating") or ""
    return ratings


def load_records(path: str) -> dict[str, dict]:
    records = {}
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        if line.strip():
            record = json.loads(line)
            records[record["query"]] = record
    return records


def entry_text(item: dict) -> str:
    """前 5 位引用条目：文档名 + chunk 定位 + 是否被判支持。"""
    lines = []
    for i, entry in enumerate(item.get("entries") or [], 1):
        mark = "支持" if entry.get("supported") else "不支持"
        chunk = entry.get("chunk_id") or entry.get("chunk_index") or "-"
        lines.append(f"{i}. {entry.get('doc') or '（无文档名）'} · {chunk} · {mark}")
    return "\n".join(lines) or "（无引用标注）"


def verdict(item: dict, record: dict) -> tuple[str, str, str]:
    """判定一格：返回（引用是否正确、判定说明、底色类别 hit/miss/na）。

    run_status 非 completed 的题先判为未作答：这类题的服务端 run 是反问澄清，
    落在记录里的正文可能是同一线程上一轮的残留，不能当作本题的答案来判引用。
    """
    if record.get("run_status") != "completed":
        return "不适用", "系统本轮未直接作答（反问澄清），无引用标注可判", "na"
    if not item.get("entries"):
        if item.get("unresolved_citation"):
            return "不适用", "答案标注的来源与本轮检索片段对不上，无法归因", "na"
        return "不适用", "答案没有可解析的引用标注", "na"
    if item.get("hit"):
        entry = item["entries"][item.get("hit_position", 1) - 1]
        return "是", f"命中第 {item.get('hit_position')} 条：{entry.get('reason') or ''}", "hit"
    return "否", "；".join((e.get("reason") or "") for e in item["entries"]), "miss"


def main() -> int:
    args = parse_args()
    records = load_records(args.results)
    ratings = load_ratings(args.ratings)
    report = json.loads(Path(args.report).read_text(encoding="utf-8"))
    items = {item["query"]: item for item in report["items"]}

    book = Workbook()
    sheet = book.active
    sheet.title = "补测明细"
    sheet.append([name for name, _ in HEADERS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="center")
    for i, (_, width) in enumerate(HEADERS, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width

    hit_count = judged = 0
    for index, (query, record) in enumerate(records.items(), 1):
        item = items.get(query, {})
        passed, reason, tone = verdict(item, record)
        if tone != "na":
            judged += 1
            hit_count += 1 if tone == "hit" else 0
        sheet.append(
            [
                index,
                record.get("section") or "",
                query,
                ratings.get(query) or "",
                record.get("gold_answer") or "",
                "" if tone == "na" and record.get("run_status") != "completed" else (record.get("agent_answer") or ""),
                entry_text(item),
                passed,
                reason,
            ]
        )
        row = sheet.max_row
        for col in (3, 5, 6, 7, 9):
            sheet.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.cell(row=row, column=8).fill = {"hit": HIT_FILL, "miss": MISS_FILL}.get(tone, NA_FILL)
        sheet.row_dimensions[row].height = 90

    sheet.freeze_panes = "A2"
    summary = (
        f"合计：补测 {len(records)} 题，其中含可归因引用标注、可判的 {judged} 题，"
        f"引用正确 {hit_count} 题，引用准确率 {100.0 * hit_count / judged:.1f}%；"
        f"其余 {len(records) - judged} 题（未直接作答 / 无引用标注 / 标注对不上检索片段）不计入分母。"
        if judged
        else f"合计：补测 {len(records)} 题，均无可归因的引用标注。"
    )
    sheet.append([summary])

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    print(f"补测 {len(records)} 题，可判 {judged} 题，引用正确 {hit_count} 题 -> {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
