#!/usr/bin/env python3
"""从《二期准确率测试.xlsx》抽出二期三批题，落成可复跑的补测题集。

二期 = MCX 30 + 终端-cat1 33 + 终端-安卓 102 = 165 题，与一期 342 题（运营平台/调度台/
MDM/miniserver/定位产品）在域上互不重叠，合计 507 = v1 摸底全量。MCX 那 30 题在一期表里
另评过一轮（旧答案、判定全「对」），与本批有 5 道结论不一致，故整批重跑、以本轮为准。

用法：
    python backend/scripts/eval_datasets/build_phase2_testset.py \
        --workbook docs/二期准确率测试.xlsx \
        --output   backend/scripts/eval_datasets/phase2_165.jsonl
"""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

SHEETS = ["MCX", "CAT1", "安卓"]
HEADER_ROW = 2


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="构建二期 165 题补测题集")
    parser.add_argument("--workbook", required=True, help="《二期准确率测试.xlsx》")
    parser.add_argument("--output", required=True, help="输出 JSONL 路径")
    parser.add_argument("--verdict", help="可选，只保留甲方判定为该值的题（如 拒答）")
    return parser.parse_args()


def load_sheet(path: str, name: str) -> list[dict]:
    """一张二期分页 → 逐行 dict（键为第 2 行表头）；第 1 行是「分区(口径) 共 N 题」标题。"""
    sheet = load_workbook(path, data_only=True)[name]
    header = [sheet.cell(row=HEADER_ROW, column=c).value for c in range(1, sheet.max_column + 1)]
    return [
        {header[c - 1]: sheet.cell(row=row, column=c).value for c in range(1, sheet.max_column + 1)}
        for row in range(HEADER_ROW + 1, sheet.max_row + 1)
    ]


def main() -> int:
    args = parse_args()
    records: list[dict] = []
    for name in SHEETS:
        rows = [row for row in load_sheet(args.workbook, name) if str(row.get("问题") or "").strip()]
        if args.verdict:
            rows = [row for row in rows if str(row.get("判定") or "").strip() == args.verdict]
        records.extend(
            {
                "index": len(records) + 1,
                "section": str(row.get("分区") or name).strip(),
                "question_no": row.get("题号"),
                "query": str(row.get("问题") or "").strip(),
                "gold_answer": str(row.get("标准答案") or "").strip(),
                "aip_verdict": str(row.get("判定") or "").strip(),
                "aip_answer": str(row.get("系统答案") or "").strip(),
            }
            for row in rows
        )

    output = Path(args.output)
    output.write_text(
        "".join(json.dumps(record, ensure_ascii=False) + "\n" for record in records), encoding="utf-8"
    )
    by_section = Counter(record["section"] for record in records)
    by_verdict = Counter(record["aip_verdict"] for record in records)
    empty_gold = sum(1 for record in records if not record["gold_answer"])
    print(f"二期题 {len(records)} 道 → {output}")
    print(f"  分区: {dict(by_section)}")
    print(f"  甲方判定: {dict(by_verdict)}；标准答案为空 {empty_gold} 道")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
