#!/usr/bin/env python3
"""从 v1 摸底产物里筛出「拒答缺口」题，落成可复跑的补测题集。

来源：《客服知识库-v1摸底507题-端到端输出.xlsx》的《明细》页，该页逐题标了分类，
`refusal_gap(拒答缺口)` 即系统当时直接拒答（答案就是那句「抱歉，在现有知识库中未找到相关依据。」）
的题。v1 摸底跑在客服知识库入库之前，所以这一批当时全数拒答——正好当作补测的「补测前」基线。

用法：
    python backend/scripts/eval_datasets/build_refusal_testset.py \
        --workbook backend/scripts/eval_datasets/final/客服知识库-v1摸底507题-端到端输出.xlsx \
        --output   backend/scripts/eval_datasets/refusal_gap148.jsonl
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

SHEET = "明细"
REFUSAL_CATEGORY = "refusal_gap(拒答缺口)"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="构建拒答题补测题集")
    parser.add_argument("--workbook", required=True, help="v1 摸底 507 题工作簿")
    parser.add_argument("--output", required=True, help="输出 JSONL 路径")
    parser.add_argument("--sections", help="可选，逗号分隔的分区白名单（如 运营平台,调度台,MDM,miniserver,定位产品）")
    return parser.parse_args()


def load_rows(path: str) -> list[dict]:
    """《明细》页 → 逐行 dict（键为表头）。"""
    sheet = load_workbook(path, data_only=True)[SHEET]
    header = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
    return [
        {header[c - 1]: sheet.cell(row=row, column=c).value for c in range(1, sheet.max_column + 1)}
        for row in range(2, sheet.max_row + 1)
    ]


def main() -> int:
    args = parse_args()
    refused = [row for row in load_rows(args.workbook) if str(row.get("分类") or "").strip() == REFUSAL_CATEGORY]
    if args.sections:
        wanted = {name.strip() for name in args.sections.split(",") if name.strip()}
        refused = [row for row in refused if str(row.get("分区") or "").strip() in wanted]
    records = [
        {
            "index": index,
            "section": str(row.get("分区") or "").strip(),
            "question_no": row.get("题号"),
            "query": str(row.get("问题") or "").strip(),
            "gold_answer": str(row.get("甲方标准答案") or "").strip(),
            "v1_answer": str(row.get("系统答案(v1摸底)") or "").strip(),
        }
        for index, row in enumerate(refused, 1)
    ]

    output = Path(args.output)
    output.write_text(
        "".join(json.dumps(record, ensure_ascii=False) + "\n" for record in records), encoding="utf-8"
    )
    sections = {record["section"] for record in records}
    print(f"拒答题 {len(records)} 道（{len(sections)} 个分区）→ {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
