#!/usr/bin/env python3
"""把拒答题补测结果导出成 Excel。

题集是甲方标为「拒答」的题（一期五域取 v1 摸底产物、二期三域取《二期准确率测试.xlsx》），
补测重跑后按定稿口径评分，这里把「补测前系统答案 / 补测后系统答案 / 分类 / 得分」逐题并列。

输入：
  - 补测跑批产物（含 query/section/phase/gold_answer/agent_answer/prior_answer）
  - 评分子报告 reports/kefu_facts_<name>.json（含 cls/gate/score/关键事实命中）

用法：
    python /app/scripts/eval_datasets/export_refusal_xlsx.py \
        --e2e    /app/scripts/eval_datasets/refusal144_e2e.jsonl \
        --report /app/scripts/eval_datasets/reports/kefu_facts_refusal144_20260929.json \
        --output /app/scripts/eval_datasets/reports/拒答题补测结果.xlsx
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    ("序号", 6),
    ("分期", 8),
    ("分区", 12),
    ("问题", 46),
    ("甲方标准答案", 46),
    ("补测前系统答案", 40),
    ("补测后系统答案", 70),
    ("分类", 12),
    ("答案正确性", 10),
    ("关键事实命中", 12),
    ("未命中的关键事实", 44),
    ("硬门槛判定依据", 40),
]

CLS_LABEL = {
    "answered": "实质作答",
    "refusal_gap": "缺口拒答",
    "clarify_missing": "反问澄清",
    "judge_error": "评测失败",
    "e2e_error": "链路异常",
}

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
ANSWERED_FILL = PatternFill("solid", fgColor="E2EFDA")
GAP_FILL = PatternFill("solid", fgColor="FCE4E4")
CLARIFY_FILL = PatternFill("solid", fgColor="FFF2CC")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="拒答题补测结果导出 Excel")
    parser.add_argument("--e2e", required=True, help="补测跑批产物 JSONL")
    parser.add_argument("--report", required=True, help="补测评分报告 JSON")
    parser.add_argument("--output", required=True, help="输出 xlsx 路径")
    return parser.parse_args()


def load_records(path: str) -> list[dict]:
    return [json.loads(line) for line in Path(path).read_text(encoding="utf-8").splitlines() if line.strip()]


def fact_cell(item: dict) -> str:
    """关键事实命中：`命中/总数`，总数为 0 时留空（未送判）。"""
    key = item.get("key_facts") or {}
    return f"{key.get('hit', 0)}/{key.get('total', 0)}" if key.get("total") else ""


def score_cell(item: dict) -> str:
    """答案正确性：只对实质作答给分，其余留空——拒答/澄清不计分，不拿 0 分冒充答错。"""
    return f"{item['score']:.2f}" if item.get("cls") == "answered" and item.get("score") is not None else ""


def summary_text(items: list[dict]) -> str:
    counts = {}
    for item in items:
        counts[item["cls"]] = counts.get(item["cls"], 0) + 1
    answered = [i for i in items if i["cls"] == "answered"]
    mean = sum(i["score"] for i in answered) / len(answered) if answered else 0.0
    failed = counts.get("judge_error", 0) + counts.get("e2e_error", 0)
    return (
        f"合计：拒答题补测 {len(items)} 题，补测前全部拒答。补测后 "
        f"实质作答 {counts.get('answered', 0)} 题、缺口拒答 {counts.get('refusal_gap', 0)} 题、"
        f"反问澄清 {counts.get('clarify_missing', 0)} 题、评测失败 {failed} 题；"
        f"答案正确性 {100.0 * mean:.1f}%（实质作答口径，拒答/澄清不计分）。"
    )


def main() -> int:
    args = parse_args()
    records = load_records(args.e2e)
    items = {item["query"]: item for item in json.loads(Path(args.report).read_text(encoding="utf-8"))["items"]}

    book = Workbook()
    sheet = book.active
    sheet.title = "拒答补测明细"
    sheet.append([name for name, _ in HEADERS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    for i, (_, width) in enumerate(HEADERS, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width

    rows = [items.get(r["query"], {}) for r in records]
    for index, (record, item) in enumerate(zip(records, rows), 1):
        sheet.append(
            [
                index,
                record.get("phase") or "",
                record.get("section") or "",
                record.get("query") or "",
                record.get("gold_answer") or "",
                record.get("prior_answer") or "",
                record.get("agent_answer") or "",
                CLS_LABEL.get(item.get("cls"), item.get("cls") or ""),
                score_cell(item),
                fact_cell(item),
                "\n".join(item.get("key_missed") or []),
                item.get("gate_reason") or "",
            ]
        )
        row = sheet.max_row
        for col in (4, 5, 6, 7, 11, 12):
            sheet.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.cell(row=row, column=8).fill = {
            "answered": ANSWERED_FILL,
            "refusal_gap": GAP_FILL,
            "clarify_missing": CLARIFY_FILL,
        }.get(item.get("cls"), GAP_FILL)
        sheet.row_dimensions[row].height = 100

    sheet.freeze_panes = "A2"
    sheet.append([summary_text(rows)])

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    print(summary_text(rows))
    print(f"-> {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
