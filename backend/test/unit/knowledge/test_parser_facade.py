from __future__ import annotations

import asyncio
import base64
import time
from pathlib import Path
from types import SimpleNamespace

import fitz
import pandas as pd
import pytest
import yuxi.knowledge.parser.unified as parser_unified
from docx import Document
from PIL import Image

from yuxi.knowledge.parser.factory import DocumentProcessorFactory
from yuxi.knowledge.parser.unified import Parser

DATA_DIR = Path(__file__).resolve().parents[2] / "data"


def _build_pdf(file_path: Path, text: str) -> None:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    doc.save(str(file_path))
    doc.close()


def _build_docx(file_path: Path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(str(file_path))


def _build_png(file_path: Path) -> None:
    image = Image.new("RGB", (120, 80), "white")
    image.save(str(file_path))


def test_parser_parse_pdf_file_returns_markdown_text(tmp_path: Path):
    file_path = tmp_path / "parser_test.pdf"
    _build_pdf(file_path, "Parser PDF content")

    markdown = Parser.parse(str(file_path), params={"ocr_engine": "disable"})

    assert isinstance(markdown, str)
    assert "Parser" in markdown
    assert "content" in markdown
    assert len(markdown.strip()) > 0


def test_parser_parse_docx_file_returns_markdown_text(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    file_path = tmp_path / "parser_test.docx"
    _build_docx(file_path, "Parser DOCX content")

    # 避免测试依赖 docling 行为，直接验证统一 parser 可回退到 python-docx。
    def _raise_docling_error(*args, **kwargs):
        raise RuntimeError("force fallback to python-docx")

    monkeypatch.setattr(parser_unified, "_convert_with_docling", _raise_docling_error)

    markdown = Parser.parse(str(file_path))

    assert isinstance(markdown, str)
    assert "Parser DOCX content" in markdown
    assert len(markdown.strip()) > 0


def test_parser_parse_html_flattens_nested_table(tmp_path: Path):
    file_path = tmp_path / "nested.html"
    file_path.write_text(
        """
        <table>
          <tr><th>Outer</th><th>Value</th></tr>
          <tr><td>Before<table><tr><th>Inner</th></tr><tr><td>Nested</td></tr></table>After</td><td>Peer</td></tr>
        </table>
        """,
        encoding="utf-8",
    )

    markdown = Parser.parse(str(file_path))

    assert "Outer" in markdown
    assert "Peer" in markdown
    assert markdown.count("Inner") == 1
    assert markdown.count("Nested") == 1


def test_parser_docx_fallback_preserves_nested_table_text(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    file_path = tmp_path / "nested.docx"
    document = Document()
    outer = document.add_table(rows=2, cols=2)
    outer.cell(0, 0).text = "Outer"
    outer.cell(0, 1).text = "Value"
    parent = outer.cell(1, 0)
    parent.text = "Before"
    nested = parent.add_table(rows=2, cols=1)
    nested.cell(0, 0).text = "Inner"
    nested.cell(1, 0).text = "Nested"
    outer.cell(1, 1).text = "Peer"
    document.save(str(file_path))

    monkeypatch.setattr(
        parser_unified,
        "_convert_with_docling",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("force fallback")),
    )

    markdown = Parser.parse(str(file_path))

    assert "| Outer | Value |" in markdown
    assert "Peer" in markdown
    assert markdown.count("Inner") == 1
    assert markdown.count("Nested") == 1


def test_convert_csv_to_markdown_preserves_column_dtypes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    file_path = tmp_path / "parser_test.csv"
    file_path.write_text("id,score\n9007199254740993,2.5\n", encoding="utf-8")
    captured_dtypes: list[dict[str, object]] = []
    original_to_markdown = pd.DataFrame.to_markdown

    def _capture_dtypes(dataframe: pd.DataFrame, *args, **kwargs) -> str:
        captured_dtypes.append(dataframe.dtypes.to_dict())
        return original_to_markdown(dataframe, *args, **kwargs)

    monkeypatch.setattr(pd.DataFrame, "to_markdown", _capture_dtypes)

    markdown = parser_unified._convert_csv_to_markdown(file_path)

    assert markdown
    assert str(captured_dtypes[0]["id"]) == "int64"


def test_convert_with_docling_reinserts_image_links_in_document_order(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    file_path = tmp_path / "parser_test.docx"
    file_path.write_bytes(b"fake docx")
    first_image = base64.b64encode(b"first image").decode()
    second_image = base64.b64encode(b"second image").decode()
    converted_markdown = "before\n<!-- image -->\nremote\n<!-- image -->\nbetween\n<!-- image -->\nafter"
    fake_doc = SimpleNamespace(
        pictures=[
            SimpleNamespace(image=SimpleNamespace(uri=f"data:image/png;base64,{first_image}")),
            SimpleNamespace(image=SimpleNamespace(uri="https://example.test/remote.png")),
            SimpleNamespace(image=SimpleNamespace(uri=f"data:image/png;base64,{second_image}")),
        ],
        export_to_markdown=lambda **_kwargs: converted_markdown,
        tables=[],
    )
    fake_result = SimpleNamespace(status=SimpleNamespace(name="SUCCESS"), document=fake_doc)
    uploaded_images: list[bytes] = []

    class FakeConverter:
        def convert(self, path: Path):
            assert path == file_path
            return fake_result

    def _fake_upload_image_to_minio(image_data, filename, bucket_name, object_prefix):
        uploaded_images.append(image_data)
        return f"https://example.test/{len(uploaded_images)}.png"

    monkeypatch.setattr(parser_unified, "_get_docling_converter", lambda: FakeConverter())
    monkeypatch.setattr(parser_unified, "_upload_image_to_minio", _fake_upload_image_to_minio)
    image_timestamps = iter([1.0, 2.0])
    monkeypatch.setattr(parser_unified.time, "time", lambda: next(image_timestamps))

    markdown = parser_unified._convert_with_docling(file_path)

    assert uploaded_images == [b"first image", b"second image"]
    assert markdown == (
        "before\n"
        "![image_1000000.png](https://example.test/1.png)\n"
        "remote\n"
        "\n"
        "between\n"
        "![image_2000000.png](https://example.test/2.png)\n"
        "after"
    )


def test_convert_with_docling_keeps_image_placeholder_when_upload_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    file_path = tmp_path / "parser_test.docx"
    file_path.write_bytes(b"fake docx")
    image = base64.b64encode(b"image data").decode()
    fake_doc = SimpleNamespace(
        pictures=[SimpleNamespace(image=SimpleNamespace(uri=f"data:image/png;base64,{image}"))],
        export_to_markdown=lambda **_kwargs: "before\n<!-- image -->\nafter",
        tables=[],
    )
    fake_result = SimpleNamespace(status=SimpleNamespace(name="SUCCESS"), document=fake_doc)

    class FakeConverter:
        def convert(self, path: Path):
            assert path == file_path
            return fake_result

    def _raise_upload_error(*args, **kwargs):
        raise RuntimeError("upload failed")

    monkeypatch.setattr(parser_unified, "_get_docling_converter", lambda: FakeConverter())
    monkeypatch.setattr(parser_unified, "_upload_image_to_minio", _raise_upload_error)
    monkeypatch.setattr(parser_unified.time, "time", lambda: 1.0)

    markdown = parser_unified._convert_with_docling(file_path)

    assert markdown == "before\n[图片: image_1000000.png]\nafter"


def test_collapse_horizontal_merges_only_shrinks_column_span() -> None:
    def _cell(start_row: int, start_col: int, end_row: int, end_col: int) -> SimpleNamespace:
        return SimpleNamespace(
            start_row_offset_idx=start_row,
            start_col_offset_idx=start_col,
            end_row_offset_idx=end_row,
            end_col_offset_idx=end_col,
        )

    merged_title = _cell(0, 0, 1, 8)  # Excel 整行合并标题 A1:H1
    merged_group = _cell(4, 0, 6, 2)  # 竖向+横向同时合并
    plain = _cell(1, 0, 2, 1)
    fake_table = SimpleNamespace(data=SimpleNamespace(table_cells=[merged_title, plain, merged_group]))
    fake_doc = SimpleNamespace(tables=[fake_table])

    parser_unified._collapse_horizontal_merges(fake_doc)

    assert (merged_title.start_col_offset_idx, merged_title.end_col_offset_idx) == (0, 1)
    assert (merged_group.start_col_offset_idx, merged_group.end_col_offset_idx) == (0, 1)
    assert (merged_group.start_row_offset_idx, merged_group.end_row_offset_idx) == (4, 6)
    assert (plain.start_col_offset_idx, plain.end_col_offset_idx) == (0, 1)


def test_parser_parse_png_file_returns_markdown_text_with_mocked_ocr(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    file_path = tmp_path / "parser_test.png"
    _build_png(file_path)

    async def _fake_parse_image_async(file, params=None):
        return "Parser PNG content"

    monkeypatch.setattr(parser_unified, "parse_image_async", _fake_parse_image_async)

    markdown = Parser.parse(str(file_path), params={"ocr_engine": "rapid_ocr"})

    assert isinstance(markdown, str)
    assert "Parser PNG content" in markdown
    assert len(markdown.strip()) > 0


def test_parse_image_uses_ocr_engine_config(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    file_path = tmp_path / "parser_test.png"
    _build_png(file_path)
    captured = {}

    def _fake_process_file(processor_type, file, params=None):
        captured["processor_type"] = processor_type
        captured["file"] = file
        captured["params"] = params
        return "OCR content"

    monkeypatch.setattr(DocumentProcessorFactory, "process_file", _fake_process_file)

    result = parser_unified.parse_image(
        str(file_path),
        params={
            "ocr_engine": "mineru_ocr",
            "backend": "old-backend",
            "ocr_engine_config": {"backend": "pipeline", "formula_enable": False},
        },
    )

    assert result == "OCR content"
    assert captured["processor_type"] == "mineru_ocr"
    assert captured["file"] == str(file_path)
    assert captured["params"]["backend"] == "pipeline"
    assert captured["params"]["formula_enable"] is False


@pytest.mark.asyncio
async def test_aparse_image_prepends_published_original(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """图片输入：原图转存公开桶后置于识别文本之前（引擎只输出碎片，原图会丢）。"""
    file_path = tmp_path / "中国通信标准化协会CCSA全权会员证书.jpg"
    _build_png(file_path)
    published: list[tuple[bytes, str, str, str]] = []

    async def _fake_parse_image_async(file, params=None):
        return "善理通益信息科技（深圳）有限公司"

    def _fake_upload_image_to_minio(image_data, filename, bucket_name, object_prefix):
        published.append((image_data, filename, bucket_name, object_prefix))
        return "http://localhost:9000/public/kb_x/kb-images/1_cert.jpg"

    monkeypatch.setattr(parser_unified, "parse_image_async", _fake_parse_image_async)
    monkeypatch.setattr(parser_unified, "_upload_image_to_minio", _fake_upload_image_to_minio)

    markdown = await Parser.aparse(str(file_path), params={"ocr_engine": "rapid_ocr"})

    assert len(published) == 1
    image_data, filename, bucket_name, object_prefix = published[0]
    assert image_data == file_path.read_bytes()
    assert filename == "中国通信标准化协会CCSA全权会员证书.jpg"
    assert (bucket_name, object_prefix) == ("public", "unknown/kb-images")
    assert markdown == (
        '<div style="text-align: center;">'
        '<img src="http://localhost:9000/public/kb_x/kb-images/1_cert.jpg" '
        'alt="中国通信标准化协会CCSA全权会员证书.jpg" /></div>'
        "\n\n善理通益信息科技（深圳）有限公司"
    )


@pytest.mark.asyncio
async def test_aparse_keeps_ocr_text_when_original_upload_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    file_path = tmp_path / "parser_test.png"
    _build_png(file_path)

    async def _fake_parse_image_async(file, params=None):
        return "OCR content"

    def _raise_upload_error(*args, **kwargs):
        raise RuntimeError("upload failed")

    monkeypatch.setattr(parser_unified, "parse_image_async", _fake_parse_image_async)
    monkeypatch.setattr(parser_unified, "_upload_image_to_minio", _raise_upload_error)

    assert await Parser.aparse(str(file_path), params={"ocr_engine": "rapid_ocr"}) == "OCR content"


def test_parse_image_ignores_enable_ocr(tmp_path: Path) -> None:
    file_path = tmp_path / "parser_test.png"
    _build_png(file_path)

    with pytest.raises(ValueError, match="必须启用OCR"):
        parser_unified.parse_image(str(file_path), params={"ocr_engine": "disable", "enable_ocr": "rapid_ocr"})


@pytest.mark.asyncio
async def test_parser_aparse_pdf_file_returns_markdown_text(tmp_path: Path):
    file_path = tmp_path / "parser_test_async.pdf"
    _build_pdf(file_path, "Async Parser PDF content")

    markdown = await Parser.aparse(str(file_path), params={"ocr_engine": "disable"})

    assert isinstance(markdown, str)
    assert "Async" in markdown
    assert "content" in markdown
    assert len(markdown.strip()) > 0


@pytest.mark.asyncio
async def test_parser_aparse_docx_does_not_block_event_loop(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    file_path = tmp_path / "parser_test_async.docx"
    file_path.write_bytes(b"fake docx")
    completion_order: list[str] = []

    def _slow_docling_conversion(*args, **kwargs) -> str:
        time.sleep(0.1)
        return "Async DOCX content"

    async def _parse_document() -> None:
        await Parser.aparse(str(file_path))
        completion_order.append("parse")

    async def _record_event_loop_progress() -> None:
        await asyncio.sleep(0.01)
        completion_order.append("event_loop")

    monkeypatch.setattr(parser_unified, "_convert_with_docling", _slow_docling_conversion)

    await asyncio.gather(_parse_document(), _record_event_loop_progress())

    assert completion_order == ["event_loop", "parse"]


def test_parse_pdf_uses_config_default_ocr_when_engine_missing(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import yuxi

    file_path = tmp_path / "parser_test.pdf"
    _build_pdf(file_path, "Parser PDF content")
    captured = {}

    def _fake_process_file(processor_type, file, params=None):
        captured["processor_type"] = processor_type
        captured["file"] = file
        captured["params"] = params
        return "default OCR content"

    monkeypatch.setattr(yuxi.config, "default_ocr_engine", "mineru_ocr")
    monkeypatch.setattr(DocumentProcessorFactory, "process_file", _fake_process_file)

    result = parser_unified.parse_pdf(str(file_path), params={})

    assert result == "default OCR content"
    assert captured["processor_type"] == "mineru_ocr"
    assert captured["file"] == str(file_path)


def test_parse_pdf_keeps_explicit_disable_when_default_ocr_enabled(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import yuxi

    file_path = tmp_path / "parser_test.pdf"
    _build_pdf(file_path, "Parser PDF content")
    monkeypatch.setattr(yuxi.config, "default_ocr_engine", "mineru_ocr")

    result = parser_unified.parse_pdf(str(file_path), params={"ocr_engine": "disable"})

    assert "Parser PDF content" in result


@pytest.mark.asyncio
async def test_parser_aparse_image_file_with_mineru_when_available():
    file_path = DATA_DIR / "测试图片.png"
    assert file_path.exists(), f"测试文件不存在: {file_path}"

    health = await asyncio.to_thread(DocumentProcessorFactory.check_health, "mineru_ocr")
    if health.get("status") != "healthy":
        pytest.skip(f"mineru_ocr 不可用: {health.get('message', 'unknown')}")

    markdown = await Parser.aparse(
        str(file_path),
        params={"ocr_engine": "mineru_ocr", "backend": "pipeline"},
    )

    assert isinstance(markdown, str)
    assert len(markdown) > 100
    assert len(markdown.strip()) > 0
